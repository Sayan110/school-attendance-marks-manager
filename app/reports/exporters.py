"""Excel / CSV export for report documents."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.report_service import ReportDocument

TEAL = "0F766E"
BORDER = Side(style="thin", color="D8E0EA")
HEADER_FILL = PatternFill("solid", fgColor=TEAL)


def _autosize(ws) -> None:
    for column in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(48, max(12, length + 3))


def document_to_excel(document: ReportDocument, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = document.school
    ws["A1"].font = Font(bold=True, size=16, color=TEAL)
    ws["A2"] = document.letterhead_line
    ws["A2"].font = Font(size=9, color="5B6B7F")
    ws["A4"] = document.title
    ws["A4"].font = Font(bold=True, size=13)
    ws["A5"] = document.subtitle
    ws["A5"].font = Font(size=9, color="5B6B7F")
    row = 7
    for label, value in document.meta:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color="5B6B7F")
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1
    if document.summary:
        ws.cell(row=row, column=1, value="Summary").font = Font(bold=True, size=11, color=TEAL)
        row += 1
        for label, value in document.summary:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value).font = Font(bold=True)
            row += 1
        row += 1
    for note in document.notes:
        ws.cell(row=row, column=1, value="• " + note).font = Font(size=9, color="5B6B7F")
        row += 1
    ws.cell(row=row + 1, column=1, value=f"Generated {document.generated_at}").font = Font(
        size=8, color="5B6B7F"
    )
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60

    for index, table in enumerate(document.tables, start=1):
        sheet = wb.create_sheet((table.title or f"Table {index}")[:31])
        sheet["A1"] = table.title or document.title
        sheet["A1"].font = Font(bold=True, size=12, color=TEAL)
        sheet["A2"] = document.school
        sheet["A2"].font = Font(size=9, color="5B6B7F")
        for column, header in enumerate(table.headers, start=1):
            cell = sheet.cell(row=4, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=BORDER)
        if table.rows:
            for record in table.rows:
                sheet.append(list(record))
        else:
            sheet.cell(row=5, column=1, value=table.empty_message).font = Font(italic=True, color="5B6B7F")
        sheet.freeze_panes = sheet.cell(row=5, column=1)
        _autosize(sheet)
    wb.save(path)
    return path


def document_to_csv(document: ReportDocument, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    table = document.tables[0] if document.tables else None
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([document.school])
        writer.writerow([document.title])
        writer.writerow([document.subtitle])
        for label, value in document.meta:
            writer.writerow([label, value])
        writer.writerow([])
        if document.summary:
            for label, value in document.summary:
                writer.writerow([label, value])
            writer.writerow([])
        if table:
            writer.writerow(table.headers)
            for row in table.rows:
                writer.writerow(["" if v is None else v for v in row])
        writer.writerow([])
        for note in document.notes:
            writer.writerow([note])
        writer.writerow([f"Generated {document.generated_at}"])
    return path
