"""CSV / Excel reading, templates and tabular exports."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.validators import ValidationError, clean_text

STUDENT_HEADERS = [
    "roll_no",
    "full_name",
    "class_name",
    "section",
    "guardian_name",
    "contact_no",
    "admission_date",
    "status",
    "remarks",
]

STUDENT_HEADER_ALIASES = {
    "roll_no": ["roll_no", "roll", "roll number", "rollno", "admission_no", "admission no", "admission number"],
    "full_name": ["full_name", "name", "student name", "student_name"],
    "class_name": ["class_name", "class", "class name"],
    "section": ["section", "sec"],
    "guardian_name": ["guardian_name", "guardian", "parent", "parent name", "father name", "mother name"],
    "contact_no": ["contact_no", "contact", "phone", "mobile", "mobile no"],
    "admission_date": ["admission_date", "admission date", "date of admission"],
    "status": ["status"],
    "remarks": ["remarks", "note", "notes"],
}


def _open_workbook_styled() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    return wb


def _style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="0F766E")
    font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D8E0EA")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=thin)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.row_dimensions[row].height = 22


def _autosize(ws) -> None:
    for column in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(42, max(12, length + 3))


def read_table_file(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    if not path.exists():
        raise ValidationError(f"The file was not found: {path}")
    suffix = path.suffix.lower()
    try:
        if suffix in (".xlsx", ".xlsm"):
            frame = pd.read_excel(path, engine="openpyxl", dtype=str)
        elif suffix == ".xls":
            frame = pd.read_excel(path, dtype=str)
        elif suffix == ".csv":
            try:
                frame = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
            except UnicodeDecodeError:
                frame = pd.read_csv(path, dtype=str, encoding="cp1252")
        else:
            raise ValidationError("Use an Excel (.xlsx) or CSV (.csv) file.")
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(f"Could not read the file:\n{exc}") from exc
    frame = frame.fillna("")
    rows = []
    for record in frame.to_dict(orient="records"):
        cleaned = {
            clean_text(key): clean_text(value) if isinstance(value, str) else value
            for key, value in record.items()
        }
        if any(clean_text(value) for value in cleaned.values()):
            rows.append(cleaned)
    if not rows:
        raise ValidationError("The file has a header but no data rows.")
    return rows


def normalize_student_row(raw: Mapping[str, Any]) -> dict[str, Any]:
    lowered = {str(key).strip().lower(): value for key, value in raw.items()}
    result: dict[str, Any] = {}
    for target, aliases in STUDENT_HEADER_ALIASES.items():
        for alias in aliases:
            if alias in lowered:
                result[target] = lowered[alias]
                break
    return result


def write_student_template(path: str | Path, sample_rows: Sequence[Sequence[Any]] | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(sample_rows or [])
    if path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(STUDENT_HEADERS)
            for row in rows:
                writer.writerow(row)
        return path
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(STUDENT_HEADERS)
    for row in rows:
        ws.append(list(row))
    _style_header(ws)
    _autosize(ws)
    notes = wb.create_sheet("Instructions")
    instructions = [
        ["Bulk Student Import — Instructions"],
        [""],
        ["1. Do not rename the header row in the 'Students' sheet."],
        ["2. roll_no, full_name and class_name are required."],
        ["3. class_name must be 9 or 10."],
        ["4. admission_date format: YYYY-MM-DD (for example 2026-04-05)."],
        ["5. status must be active, inactive or archived. Leave blank to use active."],
        ["6. The academic session is chosen in the import dialog, not in the file."],
        ["7. Duplicate roll numbers for the same class and session are rejected."],
        ["8. Existing students are never overwritten by an import."],
    ]
    for line in instructions:
        notes.append(line)
    notes.column_dimensions["A"].width = 95
    notes["A1"].font = Font(bold=True, size=14)
    wb.save(path)
    return path


def write_marks_template(
    path: str | Path,
    roster: Sequence[Mapping[str, Any]],
    theory_max: float | None = None,
    practical_max: float | None = None,
    project_max: float | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = ["roll_no", "full_name", "status"]
    if theory_max is not None:
        headers.append("theory")
    if practical_max is not None:
        headers.append("practical")
    if project_max is not None:
        headers.append("project")
    if not any(v is not None for v in (theory_max, practical_max, project_max)):
        headers.append("total")
    headers.append("remarks")
    wb = _open_workbook_styled()
    ws = wb.active
    ws.title = "Marks"
    ws.append(headers)
    for student in roster:
        row = [student.get("roll_no", ""), student.get("full_name", ""), "entered"]
        if theory_max is not None:
            row.append(student.get("theory") or "")
        if practical_max is not None:
            row.append(student.get("practical") or "")
        if project_max is not None:
            row.append(student.get("project") or "")
        if not any(v is not None for v in (theory_max, practical_max, project_max)):
            row.append(student.get("total") or "")
        row.append(student.get("remarks") or "")
        ws.append(row)
    _style_header(ws)
    _autosize(ws)
    guide = wb.create_sheet("Instructions")
    guide.append(["Marks Import — Instructions"])
    guide.append([""])
    guide.append(["status: entered, absent, exempted or pending."])
    guide.append(["Enter marks only for 'entered' rows; absent/exempted/pending need no marks."])
    guide.append(["Marks above the configured maximum are rejected."])
    guide.append(["A blank mark is never converted to zero."])
    guide.column_dimensions["A"].width = 90
    guide["A1"].font = Font(bold=True, size=14)
    wb.save(path)
    return path


def export_excel(
    path: str | Path,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    sheet_name: str = "Report",
    title: str | None = None,
    subtitle_lines: Sequence[str] = (),
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Report")[:31]
    cursor = 1
    if title:
        ws.cell(row=cursor, column=1, value=title).font = Font(bold=True, size=15)
        cursor += 1
    for line in subtitle_lines:
        ws.cell(row=cursor, column=1, value=line).font = Font(size=10, color="475569")
        cursor += 1
    if title or subtitle_lines:
        cursor += 1
    header_row = cursor
    ws.append([]) if False else None
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="D8E0EA"))
    for record in rows:
        ws.append(list(record))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    _autosize(ws)
    wb.save(path)
    return path


def export_workbook(path: str | Path, sheets: Sequence[dict[str, Any]]) -> Path:
    """sheets: [{name, headers, rows, title?}]"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for index, sheet in enumerate(sheets):
        ws = wb.create_sheet(str(sheet.get("name") or f"Sheet{index + 1}")[:31])
        cursor = 1
        if sheet.get("title"):
            ws.cell(row=1, column=1, value=sheet["title"]).font = Font(bold=True, size=14)
            cursor = 3
        headers = sheet.get("headers", [])
        for column, header in enumerate(headers, start=1):
            cell = ws.cell(row=cursor, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for record in sheet.get("rows", []):
            ws.append(list(record))
        ws.freeze_panes = ws.cell(row=cursor + 1, column=1)
        _autosize(ws)
    wb.save(path)
    return path


def export_csv(
    path: str | Path, headers: Sequence[str], rows: Iterable[Sequence[Any]]
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(list(headers))
        for row in rows:
            writer.writerow(["" if value is None else value for value in row])
    return path
