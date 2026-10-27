"""Import templates, file readers and tabular exports."""
from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.io_service import (
    STUDENT_HEADERS,
    export_csv,
    export_excel,
    normalize_student_row,
    read_table_file,
    write_marks_template,
    write_student_template,
)
from app.utils.validators import ValidationError


def test_student_excel_template_headers_and_instructions(tmp_path):
    path = write_student_template(tmp_path / "students.xlsx")
    assert path.exists()
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Students", "Instructions"]
    sheet = workbook["Students"]
    assert [cell.value for cell in sheet[1]] == STUDENT_HEADERS
    assert sheet.max_row == 1
    assert "Instructions" in workbook["Instructions"]["A1"].value


def test_student_template_with_sample_rows_reads_back(tmp_path):
    path = write_student_template(
        tmp_path / "students.xlsx",
        [["01", "Rituparna Das", "9", "A", "Guardian", "9800000000", "2026-04-05", "active", ""]],
    )
    rows = read_table_file(path)
    assert len(rows) == 1
    assert rows[0]["roll_no"] == "01"
    assert rows[0]["full_name"] == "Rituparna Das"
    assert rows[0]["class_name"] == "9"


def test_csv_template_roundtrip(tmp_path):
    path = write_student_template(
        tmp_path / "students.csv",
        [["02", "Anurag Bora", "10", "B", "", "", "", "active", "Transferred in"]],
    )
    with path.open(encoding="utf-8-sig", newline="") as handle:
        header = next(csv.reader(handle))
    assert header == STUDENT_HEADERS
    rows = read_table_file(path)
    assert rows[0]["remarks"] == "Transferred in"


def test_empty_template_is_rejected_with_clear_message(tmp_path):
    path = write_student_template(tmp_path / "empty.xlsx")
    with pytest.raises(ValidationError) as info:
        read_table_file(path)
    assert "no data rows" in str(info.value)


def test_missing_and_unsupported_files_are_rejected(tmp_path):
    with pytest.raises(ValidationError):
        read_table_file(tmp_path / "nope.csv")
    bad = tmp_path / "notes.txt"
    bad.write_text("roll_no,full_name\n01,Test\n", encoding="utf-8")
    with pytest.raises(ValidationError):
        read_table_file(bad)


def test_blank_rows_and_whitespace_are_skipped(tmp_path):
    path = tmp_path / "students.csv"
    path.write_text(
        "roll_no,full_name,class_name\n"
        "  01 ,  Rituparna Das ,9\n"
        ",,\n"
        "02,Anurag Bora,10\n",
        encoding="utf-8",
    )
    rows = read_table_file(path)
    assert [row["roll_no"] for row in rows] == ["01", "02"]
    assert rows[0]["full_name"].strip() == "Rituparna Das"


def test_normalize_student_row_accepts_aliases():
    row = normalize_student_row(
        {
            "Roll Number": "07",
            "Student Name": "Rituparna Das",
            "Class": "9",
            "Mobile No": "9800000000",
            "Parent Name": "Guardian",
            "Notes": "Prefers front bench",
        }
    )
    assert row["roll_no"] == "07"
    assert row["full_name"] == "Rituparna Das"
    assert row["class_name"] == "9"
    assert row["contact_no"] == "9800000000"
    assert row["guardian_name"] == "Guardian"
    assert row["remarks"] == "Prefers front bench"


def test_marks_template_follows_component_configuration(tmp_path):
    roster = [
        {"roll_no": "01", "full_name": "Rituparna Das"},
        {"roll_no": "02", "full_name": "Anurag Bora"},
    ]
    path = write_marks_template(
        tmp_path / "marks.xlsx", roster, theory_max=70, practical_max=30, project_max=None
    )
    workbook = load_workbook(path)
    sheet = workbook["Marks"]
    header = [cell.value for cell in sheet[1]]
    assert header == ["roll_no", "full_name", "status", "theory", "practical", "remarks"]
    assert sheet.max_row == 3
    assert sheet["C2"].value == "entered"
    assert "Instructions" in workbook.sheetnames


def test_marks_template_uses_total_when_no_components(tmp_path):
    path = write_marks_template(
        tmp_path / "marks.xlsx", [{"roll_no": "01", "full_name": "One"}]
    )
    header = [cell.value for cell in load_workbook(path)["Marks"][1]]
    assert header == ["roll_no", "full_name", "status", "total", "remarks"]


def test_export_excel_writes_title_and_all_rows(tmp_path):
    path = export_excel(
        tmp_path / "out.xlsx",
        ["Roll", "Name", "Marks"],
        [["01", "Rituparna Das", 88], ["02", "Anurag Bora", 30]],
        sheet_name="Class 9",
        title="Unit Test 1",
        subtitle_lines=["SHIKSHA VIKASH ENGLISH HIGH SCHOOL", "Class 9"],
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Class 9"]
    sheet = workbook["Class 9"]
    assert sheet["A1"].value == "Unit Test 1"
    assert sheet["A2"].value == "SHIKSHA VIKASH ENGLISH HIGH SCHOOL"
    assert sheet["A5"].value == "Roll"
    assert sheet["A6"].value == "01"
    assert sheet["B7"].value == "Anurag Bora"


def test_export_csv_preserves_unicode_and_blanks(tmp_path):
    path = export_csv(
        tmp_path / "out.csv", ["Roll", "Name", "Remarks"], [["01", "Rituparna Dâs", None]]
    )
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == ["Roll", "Name", "Remarks"]
    assert rows[1] == ["01", "Rituparna Dâs", ""]


def test_exports_create_parent_directories(tmp_path):
    target = tmp_path / "nested" / "deeper" / "out.csv"
    export_csv(target, ["A"], [["1"]])
    assert target.exists()
