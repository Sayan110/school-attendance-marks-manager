"""All ten reports build from saved data and export to PDF, Excel, CSV and HTML."""
from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.reports.exporters import document_to_csv, document_to_excel
from app.reports.html_reports import render_html
from app.reports.pdf_reports import render_pdf
from app.services.report_service import REPORT_TYPES, ReportService
from app.utils.validators import ValidationError

START = "2026-04-01"
END = "2027-03-31"
SCHOOL = "SHIKSHA VIKASH ENGLISH HIGH SCHOOL"


def build_all(ctx, dataset):
    session = ctx.session_id
    alpha = dataset["students"]["alpha"]
    beta = dataset["students"]["beta"]
    return {
        "daily_attendance": ctx.reports.daily_attendance(session, "9", "2026-10-05"),
        "monthly_attendance": ctx.reports.monthly_attendance(session, "9", 2026, 10),
        "student_attendance": ctx.reports.student_attendance(alpha, START, END),
        "class_attendance": ctx.reports.class_attendance(session, "9", START, END),
        "marksheet": ctx.reports.marksheet(dataset["exam_id"]),
        "student_progress": ctx.reports.student_progress(beta, START, END),
        "class_performance": ctx.reports.class_performance(session, "9", START, END),
        "support_report": ctx.reports.support_report(session, "9", START, END),
        "below_threshold": ctx.reports.below_threshold(session, None, START, END),
        "session_summary": ctx.reports.session_summary(session, START, END),
    }


def test_all_ten_report_types_are_registered(ctx):
    assert len(REPORT_TYPES) == 10
    from app.services.report_service import REPORT_LABELS

    assert set(REPORT_LABELS) == {key for key, _ in REPORT_TYPES}


def test_service_is_wired_into_context(ctx):
    assert isinstance(ctx.reports, ReportService)
    assert ctx.reports.analytics is ctx.analytics


def test_every_report_builds_with_school_and_metadata(ctx, dataset):
    documents = build_all(ctx, dataset)
    assert len(documents) == 10
    for key, document in documents.items():
        assert document.key == key
        assert document.title
        assert document.school == SCHOOL
        assert document.generated_at
        assert document.meta, f"{key} has no metadata"
        labels = [label for label, _ in document.meta]
        assert "Academic session" in labels
        assert "Subject" in labels
        assert document.subtitle


def test_daily_attendance_counts_and_rows(ctx, dataset):
    document = ctx.reports.daily_attendance(ctx.session_id, "9", "2026-10-05")
    summary = dict(document.summary)
    assert summary["Present"] == "1"
    assert summary["Absent"] == "1"
    assert summary["Leave"] == "0"
    assert summary["Unmarked"] == "0"
    assert summary["Attendance %"] == "50.00%"
    rows = {row[1]: row for row in document.tables[0].rows}
    assert rows["Alpha Sharma"][3] == "Present"
    assert rows["Beta Das"][3] == "Absent"
    assert document.landscape is False


def test_unmarked_students_are_not_reported_as_absent(ctx, dataset):
    document = ctx.reports.daily_attendance(ctx.session_id, "10", "2026-10-09")
    summary = dict(document.summary)
    assert summary["Unmarked"] == "1"
    assert summary["Absent"] == "0"
    assert summary["Attendance %"] == "—"
    assert document.tables[0].rows[0][3] == "Unmarked"


def test_monthly_register_marks_holidays_and_blanks(ctx, dataset):
    ctx.attendance.add_holiday(ctx.session_id, "2026-10-08", "Local holiday")
    document = ctx.reports.monthly_attendance(ctx.session_id, "9", 2026, 10)
    assert document.landscape is True
    table = document.tables[0]
    assert len(table.rows) == 2
    assert table.headers[2] == "1"
    assert len(table.headers) == 2 + 31 + 4
    holiday_column = table.headers.index("8")
    assert all(row[holiday_column] == "H" for row in table.rows)
    assert any("Local holiday" in note for note in document.notes)
    alpha = table.rows[0]
    summary = dict(document.summary)
    assert summary["Students"] == "2"
    assert summary["Present"] == "2"
    assert summary["Absent"] == "1"
    assert summary["Leave"] == "1"


def test_student_attendance_report_only_covers_period(ctx, dataset):
    alpha = dataset["students"]["alpha"]
    document = ctx.reports.student_attendance(alpha, "2026-10-05", "2026-10-05")
    assert len(document.tables[0].rows) == 1
    summary = dict(document.summary)
    assert summary["Present"] == "1"
    assert summary["Attendance %"] == "100.00%"


def test_marksheet_marks_pending_separately_and_ranks(ctx, dataset, add_student):
    add_student("03", "Pending Pupil", "9")
    document = ctx.reports.marksheet(dataset["exam_id"])
    rows = {row[1]: row for row in document.tables[0].rows}
    assert rows["Alpha Sharma"][-4] == "Passed"
    assert rows["Alpha Sharma"][-3] == "1"
    assert rows["Beta Das"][-4] == "Failed"
    assert rows["Pending Pupil"][-4] == "Pending"
    assert rows["Pending Pupil"][-3] == "—"
    assert rows["Pending Pupil"][2] == "—"
    summary = dict(document.summary)
    assert summary["Pass %"] == "50.00%"
    assert summary["Pending"] == "1"
    assert any("never converted to zero" in note for note in document.notes)
    assert any("not an official SEBA" in note for note in document.notes)


def test_marksheet_rejects_unknown_exam(ctx):
    with pytest.raises(ValidationError):
        ctx.reports.marksheet(9999)


def test_below_threshold_uses_configured_setting(ctx, dataset):
    document = ctx.reports.below_threshold(ctx.session_id, None, START, END)
    assert [row[1] for row in document.tables[0].rows] == ["Beta Das"]
    ctx.settings.set_many({"attendance_threshold": "101"})
    try:
        strict = ctx.reports.below_threshold(ctx.session_id, None, START, END)
        assert len(strict.tables[0].rows) == 3
    finally:
        ctx.settings.set_many({"attendance_threshold": "75"})


def test_support_report_lists_failed_results(ctx, dataset):
    document = ctx.reports.support_report(ctx.session_id, "9", START, END)
    failures = document.tables[0]
    assert [row[3] for row in failures.rows] == ["Beta Das"]
    support = {row[1] for row in document.tables[1].rows}
    assert "Beta Das" in support
    assert any("not counted as a failure" in note for note in document.notes)


def test_class_performance_bundles_exams_and_students(ctx, dataset):
    document = ctx.reports.class_performance(ctx.session_id, "9", START, END)
    assert document.landscape is True
    assert len(document.tables) == 2
    assert document.tables[0].rows[0][0] == "Unit Test 1"
    names = {row[1] for row in document.tables[1].rows}
    assert names == {"Alpha Sharma", "Beta Das"}


def test_session_summary_totals(ctx, dataset):
    document = ctx.reports.session_summary(ctx.session_id, START, END)
    summary = dict(document.summary)
    assert summary["Class 9 students"] == "2"
    assert summary["Class 10 students"] == "1"
    classes = document.tables[0].rows
    assert [row[0] for row in classes] == ["Class 9", "Class 10"]


def test_build_dispatch_rejects_unknown_key(ctx):
    with pytest.raises(ValidationError):
        ctx.reports.build("not_a_report")


def test_every_report_renders_a4_pdf(ctx, dataset, tmp_path):
    documents = build_all(ctx, dataset)
    for key, document in documents.items():
        target = tmp_path / f"{key}.pdf"
        render_pdf(document, target)
        payload = target.read_bytes()
        assert payload.startswith(b"%PDF"), key
        assert len(payload) > 1500, key


def test_reports_render_when_no_data_exists(ctx, tmp_path):
    empty = ctx.reports.daily_attendance(ctx.session_id, "9", "2026-10-05")
    assert empty.tables[0].rows == []
    assert empty.has_tables is False
    target = tmp_path / "empty.pdf"
    render_pdf(empty, target)
    assert target.stat().st_size > 1000
    html = render_html(empty)
    assert "No students are enrolled" in html


def test_html_preview_contains_school_tables_and_notes(ctx, dataset):
    document = ctx.reports.marksheet(dataset["exam_id"])
    html = render_html(document)
    assert SCHOOL in html
    assert "<table class='data'>" in html
    assert "Alpha Sharma" in html
    assert "Generated on" in html
    assert "Notes and rules" in html


def test_html_preview_escapes_unsafe_text(ctx, dataset):
    document = ctx.reports.daily_attendance(ctx.session_id, "9", "2026-10-05")
    document.notes.append("<script>alert('x')</script> & more")
    html = render_html(document)
    assert "<script>" not in html
    assert "&lt;script&gt;" in html


def test_excel_export_has_letterhead_and_one_sheet_per_table(ctx, dataset, tmp_path):
    document = ctx.reports.marksheet(dataset["exam_id"])
    target = tmp_path / "marksheet.xlsx"
    document_to_excel(document, target)
    workbook = load_workbook(target)
    assert workbook.sheetnames[0] == "Report"
    first = workbook[workbook.sheetnames[0]]
    assert first["A1"].value == SCHOOL
    assert "SEBA Board, Assam" in first["A2"].value
    assert first["A4"].value == document.title
    data_sheet = workbook[workbook.sheetnames[1]]
    assert data_sheet["A4"].value == "Roll"
    assert data_sheet.max_row >= 3


def test_csv_export_starts_with_school_and_headers(ctx, dataset, tmp_path):
    document = ctx.reports.class_attendance(ctx.session_id, "9", START, END)
    target = tmp_path / "class.csv"
    document_to_csv(document, target)
    with target.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == [SCHOOL]
    assert rows[1] == [document.title]
    header_index = next(i for i, row in enumerate(rows) if row[:2] == ["Roll", "Student name"])
    assert rows[header_index + 1][1] == "Alpha Sharma"
    assert rows[-1][0].startswith("Generated ")


def test_export_empty_table_produces_readable_file(ctx, tmp_path):
    document = ctx.reports.daily_attendance(ctx.session_id, "9", "2026-10-05")
    excel = document_to_excel(document, tmp_path / "empty.xlsx")
    workbook = load_workbook(excel)
    assert workbook[workbook.sheetnames[1]]["A5"].value == document.tables[0].empty_message
    csv_path = document_to_csv(document, tmp_path / "empty.csv")
    assert csv_path.stat().st_size > 0
